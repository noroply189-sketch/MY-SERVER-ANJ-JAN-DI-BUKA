from __future__ import annotations

import csv
import io
import json
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


OUTPUT_DIR = Path(__file__).resolve().parents[1] / "storage" / "output"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def write_text_file(filename: str, content: str) -> Path:
    path = OUTPUT_DIR / filename
    path.write_text(content, encoding="utf-8")
    return path


def write_json_file(filename: str, payload: dict) -> Path:
    path = OUTPUT_DIR / filename
    with path.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
    return path


def write_pdf_file(filename: str, text: str) -> Path:
    path = OUTPUT_DIR / filename
    pdf = canvas.Canvas(str(path), pagesize=letter)
    lines = text.splitlines()
    y = 750
    for line in lines:
        pdf.drawString(50, y, line[:120])
        y -= 18
        if y < 50:
            pdf.showPage()
            y = 750
    pdf.save()
    return path


def write_docx_file(filename: str, paragraphs: list[str]) -> Path:
    doc = Document()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    path = OUTPUT_DIR / filename
    doc.save(path)
    return path


def write_xlsx_file(filename: str, rows: list[list[object]], sheet_name: str = "Sheet1") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    path = OUTPUT_DIR / filename
    wb.save(path)
    return path


def write_csv_file(filename: str, rows: list[list[object]]) -> Path:
    path = OUTPUT_DIR / filename
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerows(rows)
    return path


def write_md_file(filename: str, content: str) -> Path:
    return write_text_file(filename, content)


def write_html_file(filename: str, content: str) -> Path:
    return write_text_file(filename, content)


def write_code_file(filename: str, code: str) -> Path:
    return write_text_file(filename, code)


def write_pptx_file(filename: str, slides: list[str]) -> Path:
    prs = Presentation()
    for slide_text in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title = slide.shapes.title if slide.shapes.title else slide.shapes.add_textbox(0.5 * 914400, 0.5 * 914400, 8 * 914400, 0.5 * 914400)
        title.text = "Slide"
        body = slide.shapes.add_textbox(0.5 * 914400, 1.2 * 914400, 8.5 * 914400, 5 * 914400)
        body.text_frame.text = slide_text
    path = OUTPUT_DIR / filename
    prs.save(path)
    return path


def make_zip_archive(filename: str, files: list[str | Path]) -> Path:
    path = OUTPUT_DIR / filename
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for file in files:
            file_path = Path(file)
            if file_path.exists():
                archive.write(file_path, arcname=file_path.name)
    return path


__all__ = [
    "OUTPUT_DIR",
    "write_text_file",
    "write_json_file",
    "write_pdf_file",
    "write_docx_file",
    "write_xlsx_file",
    "write_csv_file",
    "write_md_file",
    "write_html_file",
    "write_code_file",
    "write_pptx_file",
    "make_zip_archive",
]
