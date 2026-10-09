from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

import fitz
from docx import Document
from openpyxl import load_workbook

from utils.security import is_archive_safe, sanitize_filename
from utils.validators import validate_extension, validate_file_size

SUPPORTED_TEXT_EXTENSIONS = {
    "txt", "md", "rtf", "odt", "json", "xml", "yaml", "yml", "sql", "py", "js", "ts",
    "html", "css", "log", "csv", "pdf", "docx", "xlsx", "pptx"
}


def _read_text_file(path: str | Path) -> str:
    file_path = Path(path)
    if file_path.suffix.lower() == ".csv":
        with file_path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
            return handle.read()
    with file_path.open("r", encoding="utf-8", errors="replace") as handle:
        return handle.read()


def _read_pdf(path: str | Path) -> str:
    doc = fitz.open(path)
    text_chunks = []
    for page in doc:
        chunk = page.get_text()
        if chunk.strip():
            text_chunks.append(chunk)
    doc.close()
    return "\n\n".join(text_chunks)


def _read_docx(path: str | Path) -> str:
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)


def _read_xlsx(path: str | Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows(values_only=True):
            if any(value is not None for value in row):
                rows.append({"sheet": sheet_name, "row": list(row)})
    return rows


def read_document(path: str | Path, max_size_mb: int = 20) -> str:
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")
    if not validate_file_size(file_path, max_size_mb):
        raise ValueError(f"File exceeds the allowed size of {max_size_mb} MB.")
    ext = file_path.suffix.lower().lstrip(".")
    if ext not in SUPPORTED_TEXT_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}")

    if ext == "pdf":
        return _read_pdf(file_path)
    if ext == "docx":
        return _read_docx(file_path)
    if ext in {"xlsx", "xls"}:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        sheets = []
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            rows = []
            for row in sheet.iter_rows(values_only=True):
                if any(v is not None for v in row):
                    rows.append(list(row))
            sheets.append(f"## {sheet_name}\n" + "\n".join(", ".join(str(v) for v in row) for row in rows[:50]))
        return "\n\n".join(sheets)
    return _read_text_file(file_path)


def summarize_document(path: str | Path, max_size_mb: int = 20) -> str:
    content = read_document(path, max_size_mb=max_size_mb)
    sentences = [part.strip() for part in content.replace("\n", " ").split(".") if part.strip()]
    summary = " ".join(sentences[:10])
    return summary[:4000] if len(summary) > 4000 else summary


def extract_tables_from_csv(path: str | Path) -> list[list[str]]:
    file_path = Path(path)
    with file_path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
        reader = csv.reader(handle)
        return [row for row in reader if any(cell.strip() for cell in row)]


def read_zip_archive(path: str | Path) -> list[str]:
    file_path = Path(path)
    if not is_archive_safe(file_path):
        raise ValueError("The archive is unsafe or too large.")
    import zipfile

    entries: list[str] = []
    with zipfile.ZipFile(file_path) as archive:
        for info in archive.infolist():
            if info.is_dir():
                continue
            entries.append(info.filename)
    return entries


__all__ = [
    "read_document",
    "summarize_document",
    "extract_tables_from_csv",
    "read_zip_archive",
    "SUPPORTED_TEXT_EXTENSIONS",
]
